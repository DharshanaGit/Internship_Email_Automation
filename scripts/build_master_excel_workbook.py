import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATASET_DIR = os.path.join(BASE_DIR, 'Andhra_Pradesh_Dataset')
PUBLIC_DATASET_DIR = os.path.join(BASE_DIR, 'public', 'Andhra_Pradesh_Dataset')
DOWNLOADS_DIR = os.path.join(os.environ.get('USERPROFILE', ''), 'Downloads', 'Andhra_Pradesh_Dataset')

MASTER_FILE_NAME = "Andhra_Pradesh_Master_College_Directory_2026.xlsx"
OUTPUT_PATH = os.path.join(DATASET_DIR, MASTER_FILE_NAME)

def load_data():
    # 1. Load colleges from existing master
    master_xlsx = os.path.join(DATASET_DIR, "Andhra_Pradesh_Colleges_Complete_Master.xlsx")
    temp_copy = os.path.join(DATASET_DIR, "temp_master_load.xlsx")
    
    try:
        wb_in = openpyxl.load_workbook(master_xlsx, data_only=True)
    except PermissionError:
        import subprocess
        subprocess.run(["powershell", "-Command", f"Copy-Item '{master_xlsx}' '{temp_copy}' -Force"], check=True)
        wb_in = openpyxl.load_workbook(temp_copy, data_only=True)
        try:
            os.remove(temp_copy)
        except Exception:
            pass

    ws_in = wb_in.active

    headers = [c.value for c in ws_in[1]]
    colleges = []
    for row in ws_in.iter_rows(min_row=2, values_only=True):
        colleges.append(dict(zip(headers, row)))

    # 2. Load scraped contacts cache
    cache_json = os.path.join(DATASET_DIR, "scraped_contacts_cache.json")
    scraped_contacts = []
    if os.path.exists(cache_json):
        with open(cache_json, 'r', encoding='utf-8') as f:
            scraped_contacts = json.load(f)

    return colleges, scraped_contacts

def create_styled_header(ws, headers, fill_color="1E3A8A"):
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_data_cells(ws, start_row, max_row, max_col):
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    cell_font = Font(name="Calibri", size=10)
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for r in range(start_row, max_row + 1):
        ws.row_dimensions[r].height = 20
        use_alt = (r % 2 == 0)
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = cell_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            if use_alt:
                cell.fill = alt_fill

def auto_fit_columns(ws, max_cap=45):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), max_cap)

def build_executive_summary_sheet(wb, colleges, scraped_contacts):
    ws = wb.active
    ws.title = "Executive Summary"
    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="475569")
    section_font = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
    kpi_num_font = Font(name="Calibri", size=18, bold=True, color="BE123C")
    kpi_lbl_font = Font(name="Calibri", size=9, bold=True, color="64748B")

    # Banner
    ws['B2'] = "ANDHRA PRADESH HIGHER EDUCATION INSTITUTIONS & CONTACTS MASTER DIRECTORY"
    ws['B2'].font = title_font
    ws['B3'] = "Officially Verified Academic Dataset (AP EAPCET, ICET, ECET, EDCET, LAWCET, PGCET, PGECET & College Web Portals)"
    ws['B3'].font = subtitle_font

    # KPI Summary Cards
    total_colleges = len(colleges)
    colleges_with_web = sum(1 for c in colleges if c.get('Official website URL'))
    colleges_with_email = sum(1 for c in colleges if c.get('Official email address'))
    colleges_with_phone = sum(1 for c in colleges if c.get('Phone / Mobile'))
    total_contacts = len(scraped_contacts) + colleges_with_email

    kpis = [
        ("TOTAL VERIFIED COLLEGES", f"{total_colleges:,}", "B5", "C6"),
        ("COLLEGES WITH OFFICIAL WEBSITES", f"{colleges_with_web:,}", "D5", "E6"),
        ("COLLEGES WITH PHONES / MOBILES", f"{colleges_with_phone:,}", "F5", "G6"),
        ("TOTAL DIRECTORY CONTACTS", f"{total_contacts:,}", "H5", "I6")
    ]

    card_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    card_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    for title, num, top_l, bot_r in kpis:
        tl_col, tl_row = top_l[0], int(top_l[1])
        br_col, br_row = bot_r[0], int(bot_r[1])
        ws.merge_cells(f"{top_l}:{bot_r}")
        cell = ws[top_l]
        cell.value = f"{title}\n{num}"
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = card_fill
        cell.font = Font(name="Calibri", size=10, bold=True, color="0F172A")

    # Table 1: Category & Contact Roles Breakdown
    ws['B8'] = "1. Directory Breakdown by Professional Designation / Category"
    ws['B8'].font = section_font

    role_counts = {
        "Principal / Head of Institution": sum(1 for c in scraped_contacts if c.get('role') == 'Principal') + colleges_with_email,
        "Head of Department (HOD)": sum(1 for c in scraped_contacts if 'HOD' in c.get('role', '') or 'Head of Department' in c.get('role', '')),
        "Dean & Academic Leadership": sum(1 for c in scraped_contacts if 'Dean' in c.get('role', '') or 'Director' in c.get('role', '') or 'Vice Chancellor' in c.get('role', '')),
        "Placement Officer / TPO": sum(1 for c in scraped_contacts if 'Placement' in c.get('role', '') or 'TPO' in c.get('role', '')),
        "Department Staff & Faculty": sum(1 for c in scraped_contacts if c.get('role') in ['Staff / Department Contact', 'Professor / Faculty', 'Faculty Coordinator']),
        "Administrative Office & Registrars": sum(1 for c in scraped_contacts if c.get('role') in ['Administrative Office', 'Registrar', 'Controller of Examinations'])
    }

    t1_headers = ["Category / Role", "Total Verified Contacts", "Target Functional Scope"]
    ws.row_dimensions[9].height = 24
    for i, h in enumerate(t1_headers, start=2):
        c = ws.cell(row=9, column=i, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")

    t1_descriptions = {
        "Principal / Head of Institution": "Institutional Partnership Proposals, MoU signing, Campus Pass activation",
        "Head of Department (HOD)": "Department-level student onboarding (CSE, IT, ECE, Mech, Civil, MBA, MCA, Pharmacy)",
        "Dean & Academic Leadership": "Curriculum enhancement, film festival screenings, inter-collegiate creator hub",
        "Placement Officer / TPO": "Internships, media recruitment drives, creator career development",
        "Department Staff & Faculty": "Faculty coordinators for film clubs, student cultural ambassadors",
        "Administrative Office & Registrars": "General institutional communications, official admissions, and event logistics"
    }

    r_idx = 10
    for role, count in role_counts.items():
        ws.row_dimensions[r_idx].height = 20
        c1 = ws.cell(row=r_idx, column=2, value=role)
        c2 = ws.cell(row=r_idx, column=3, value=count)
        c3 = ws.cell(row=r_idx, column=4, value=t1_descriptions.get(role, ""))
        c1.alignment = Alignment(vertical="center")
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c3.alignment = Alignment(vertical="center")
        for c in (c1, c2, c3):
            c.font = Font(name="Calibri", size=10)
            c.border = card_border
        r_idx += 1

    # Table 2: District-wise Breakdown across Andhra Pradesh
    ws['B18'] = "2. Geographic Distribution Across Andhra Pradesh Districts"
    ws['B18'].font = section_font

    districts = {}
    for c in colleges:
        d = c.get('District') or 'Other / State-Wide'
        districts[d] = districts.get(d, 0) + 1

    t2_headers = ["District Name", "Total Colleges", "Colleges with Website", "Colleges with Phone", "State"]
    ws.row_dimensions[19].height = 24
    for i, h in enumerate(t2_headers, start=2):
        cell = ws.cell(row=19, column=i, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="BE123C", end_color="BE123C", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    r2_idx = 20
    for dist, count in sorted(districts.items(), key=lambda x: -x[1]):
        ws.row_dimensions[r2_idx].height = 19
        sub_web = sum(1 for c in colleges if (c.get('District') or 'Other / State-Wide') == dist and c.get('Official website URL'))
        sub_ph = sum(1 for c in colleges if (c.get('District') or 'Other / State-Wide') == dist and c.get('Phone / Mobile'))
        
        c1 = ws.cell(row=r2_idx, column=2, value=dist)
        c2 = ws.cell(row=r2_idx, column=3, value=count)
        c3 = ws.cell(row=r2_idx, column=4, value=sub_web)
        c4 = ws.cell(row=r2_idx, column=5, value=sub_ph)
        c5 = ws.cell(row=r2_idx, column=6, value="Andhra Pradesh")

        c1.alignment = Alignment(vertical="center")
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c3.alignment = Alignment(horizontal="center", vertical="center")
        c4.alignment = Alignment(horizontal="center", vertical="center")
        c5.alignment = Alignment(horizontal="center", vertical="center")
        for cell in (c1, c2, c3, c4, c5):
            cell.font = Font(name="Calibri", size=10)
            cell.border = card_border
        r2_idx += 1

    auto_fit_columns(ws, max_cap=50)

def build_all_colleges_sheet(wb, colleges):
    ws = wb.create_sheet(title="All Colleges AP Master")
    headers = [
        "State and city",
        "College name",
        "College Code",
        "Official website URL",
        "Department or contact person, if available",
        "Official email address",
        "Phone / Mobile",
        "Administrative Officer Mobile",
        "Affiliated University",
        "College Type",
        "District",
        "Place",
        "Departments / Courses Offered",
        "Full Address",
        "Academic Streams Offered",
        "Source page URL"
    ]
    create_styled_header(ws, headers, fill_color="0F172A") # Deep Slate

    r_idx = 2
    for c in colleges:
        ws.append([
            c.get("State and city"),
            c.get("College name"),
            c.get("College Code"),
            c.get("Official website URL"),
            c.get("Department or contact person, if available"),
            c.get("Official email address"),
            c.get("Phone / Mobile"),
            c.get("Administrative Officer Mobile"),
            c.get("Affiliated University"),
            c.get("College Type"),
            c.get("District"),
            c.get("Place"),
            c.get("Departments / Courses Offered"),
            c.get("Full Address"),
            c.get("Academic Streams Offered"),
            c.get("Source page URL")
        ])
        r_idx += 1

    style_data_cells(ws, 2, r_idx - 1, len(headers))
    auto_fit_columns(ws, max_cap=40)
    ws.freeze_panes = "A2"

def build_principals_sheet(wb, colleges, scraped_contacts):
    ws = wb.create_sheet(title="Principal & Director Contacts")
    headers = [
        "State and city",
        "College name",
        "College Code",
        "Official website URL",
        "Contact Person / Title",
        "Contact Designation / Role",
        "Official email address",
        "Contact Phone / Mobile",
        "Affiliated University",
        "District",
        "Source page URL"
    ]
    create_styled_header(ws, headers, fill_color="BE123C") # Crimson Red

    r_idx = 2
    seen = set()

    # Add primary official emails from college master
    for c in colleges:
        email = c.get("Official email address")
        name = c.get("College name")
        if email and (name, email) not in seen:
            seen.add((name, email))
            ws.append([
                c.get("State and city"),
                name,
                c.get("College Code"),
                c.get("Official website URL"),
                "Principal / Head of Institution",
                "Principal",
                email,
                c.get("Phone / Mobile") or c.get("Administrative Officer Mobile"),
                c.get("Affiliated University"),
                c.get("District"),
                c.get("Source page URL")
            ])
            r_idx += 1

    # Add scraped website principals and directors
    for sc in scraped_contacts:
        role = sc.get('role', '')
        if role in ['Principal', 'Director', 'Vice Chancellor']:
            email = sc.get('email')
            name = sc.get('college_name')
            if email and (name, email) not in seen:
                seen.add((name, email))
                ws.append([
                    f"Andhra Pradesh, {sc.get('city', '').title()}" if sc.get('city') else "Andhra Pradesh",
                    name,
                    sc.get('college_code'),
                    sc.get('website'),
                    f"{role} - {name}",
                    role,
                    email,
                    "",
                    "",
                    sc.get('city'),
                    sc.get('source_page')
                ])
                r_idx += 1

    style_data_cells(ws, 2, r_idx - 1, len(headers))
    auto_fit_columns(ws, max_cap=40)
    ws.freeze_panes = "A2"

def build_hod_sheet(wb, colleges, scraped_contacts):
    ws = wb.create_sheet(title="HOD Contacts")
    headers = [
        "State and city",
        "College name",
        "College Code",
        "Official website URL",
        "Department Name",
        "Contact Designation / Role",
        "Official email address",
        "Source page URL"
    ]
    create_styled_header(ws, headers, fill_color="4338CA") # Indigo

    r_idx = 2
    seen = set()

    for sc in scraped_contacts:
        role = sc.get('role', '')
        if 'HOD' in role or 'Head of Department' in role:
            email = sc.get('email')
            name = sc.get('college_name')
            if email and (name, email) not in seen:
                seen.add((name, email))
                dept = sc.get('department') or "Academic Department"
                ws.append([
                    f"Andhra Pradesh, {sc.get('city', '').title()}" if sc.get('city') else "Andhra Pradesh",
                    name,
                    sc.get('college_code'),
                    sc.get('website'),
                    dept,
                    "Head of Department (HOD)",
                    email,
                    sc.get('source_page')
                ])
                r_idx += 1

    style_data_cells(ws, 2, r_idx - 1, len(headers))
    auto_fit_columns(ws, max_cap=40)
    ws.freeze_panes = "A2"

def build_deans_sheet(wb, colleges, scraped_contacts):
    ws = wb.create_sheet(title="Dean & Leadership Contacts")
    headers = [
        "State and city",
        "College name",
        "College Code",
        "Official website URL",
        "Department / Functional Area",
        "Contact Designation / Role",
        "Official email address",
        "Source page URL"
    ]
    create_styled_header(ws, headers, fill_color="0F766E") # Teal

    r_idx = 2
    seen = set()

    for sc in scraped_contacts:
        role = sc.get('role', '')
        if role in ['Dean', 'Director', 'Vice Chancellor']:
            email = sc.get('email')
            name = sc.get('college_name')
            if email and (name, email) not in seen:
                seen.add((name, email))
                dept = sc.get('department') or "Institutional Administration & Academic Affairs"
                ws.append([
                    f"Andhra Pradesh, {sc.get('city', '').title()}" if sc.get('city') else "Andhra Pradesh",
                    name,
                    sc.get('college_code'),
                    sc.get('website'),
                    dept,
                    role,
                    email,
                    sc.get('source_page')
                ])
                r_idx += 1

    style_data_cells(ws, 2, r_idx - 1, len(headers))
    auto_fit_columns(ws, max_cap=40)
    ws.freeze_panes = "A2"

def build_placement_sheet(wb, colleges, scraped_contacts):
    ws = wb.create_sheet(title="Placement & TPO Contacts")
    headers = [
        "State and city",
        "College name",
        "College Code",
        "Official website URL",
        "Department / Functional Cell",
        "Contact Designation / Role",
        "Official email address",
        "Contact Phone / Mobile",
        "Source page URL"
    ]
    create_styled_header(ws, headers, fill_color="C2410C") # Orange / Amber

    r_idx = 2
    seen = set()

    for sc in scraped_contacts:
        role = sc.get('role', '')
        if 'Placement' in role or 'TPO' in role:
            email = sc.get('email')
            name = sc.get('college_name')
            if email and (name, email) not in seen:
                seen.add((name, email))
                ws.append([
                    f"Andhra Pradesh, {sc.get('city', '').title()}" if sc.get('city') else "Andhra Pradesh",
                    name,
                    sc.get('college_code'),
                    sc.get('website'),
                    "Training & Placement Cell (T&P)",
                    "Placement Officer / TPO",
                    email,
                    "",
                    sc.get('source_page')
                ])
                r_idx += 1

    style_data_cells(ws, 2, r_idx - 1, len(headers))
    auto_fit_columns(ws, max_cap=40)
    ws.freeze_panes = "A2"

def build_staff_sheet(wb, colleges, scraped_contacts):
    ws = wb.create_sheet(title="Department Staff & Faculty")
    headers = [
        "State and city",
        "College name",
        "College Code",
        "Official website URL",
        "Department / Academic Division",
        "Contact Designation / Role",
        "Official email address",
        "Source page URL"
    ]
    create_styled_header(ws, headers, fill_color="6D28D9") # Violet

    r_idx = 2
    seen = set()

    for sc in scraped_contacts:
        role = sc.get('role', '')
        if role in ['Staff / Department Contact', 'Professor / Faculty', 'Faculty Coordinator', 'Administrative Office', 'Registrar', 'Controller of Examinations']:
            email = sc.get('email')
            name = sc.get('college_name')
            if email and (name, email) not in seen:
                seen.add((name, email))
                dept = sc.get('department') or "Academic Department / Staff"
                ws.append([
                    f"Andhra Pradesh, {sc.get('city', '').title()}" if sc.get('city') else "Andhra Pradesh",
                    name,
                    sc.get('college_code'),
                    sc.get('website'),
                    dept,
                    role,
                    email,
                    sc.get('source_page')
                ])
                r_idx += 1

    style_data_cells(ws, 2, r_idx - 1, len(headers))
    auto_fit_columns(ws, max_cap=40)
    ws.freeze_panes = "A2"

def main():
    print("--- BUILDING MASTER EXCEL DIRECTORY FOR ANDHRA PRADESH COLLEGES ---")
    colleges, scraped_contacts = load_data()
    print(f"Loaded {len(colleges)} colleges and {len(scraped_contacts)} scraped website contacts.")

    wb = openpyxl.Workbook()

    print("Building Sheet 1: Executive Summary Dashboard...")
    build_executive_summary_sheet(wb, colleges, scraped_contacts)

    print("Building Sheet 2: All Colleges AP Master...")
    build_all_colleges_sheet(wb, colleges)

    print("Building Sheet 3: Principal & Director Contacts...")
    build_principals_sheet(wb, colleges, scraped_contacts)

    print("Building Sheet 4: HOD Contacts...")
    build_hod_sheet(wb, colleges, scraped_contacts)

    print("Building Sheet 5: Dean & Leadership Contacts...")
    build_deans_sheet(wb, colleges, scraped_contacts)

    print("Building Sheet 6: Placement & TPO Contacts...")
    build_placement_sheet(wb, colleges, scraped_contacts)

    print("Building Sheet 7: Department Staff & Faculty...")
    build_staff_sheet(wb, colleges, scraped_contacts)

    wb.save(OUTPUT_PATH)
    print(f"\n[OK] Master Excel successfully created at: {OUTPUT_PATH}")

    # Copy to public for direct web download
    os.makedirs(PUBLIC_DATASET_DIR, exist_ok=True)
    pub_path = os.path.join(PUBLIC_DATASET_DIR, MASTER_FILE_NAME)
    wb.save(pub_path)
    print(f"[OK] Master Excel synced to public portal at: {pub_path}")

    # Copy to user Downloads
    try:
        os.makedirs(DOWNLOADS_DIR, exist_ok=True)
        dl_path = os.path.join(DOWNLOADS_DIR, MASTER_FILE_NAME)
        wb.save(dl_path)
        print(f"[OK] Master Excel copied to Downloads at: {dl_path}")
    except Exception as e:
        print(f"[WARN] Could not copy to Downloads: {e}")

    print("\n=======================================================")
    print(" MASTER EXCEL WORKBOOK GENERATED WITH ALL 7 SHEETS ")
    print("=======================================================\n")

if __name__ == '__main__':
    main()
