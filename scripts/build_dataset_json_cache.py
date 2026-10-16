import openpyxl
import json
import os

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATASET_DIR = os.path.join(BASE_DIR, 'Andhra_Pradesh_Dataset')
PUBLIC_DATASET_DIR = os.path.join(BASE_DIR, 'public', 'Andhra_Pradesh_Dataset')

xlsx_path = os.path.join(DATASET_DIR, 'Andhra_Pradesh_Master_College_Directory_2026_Updated.xlsx')
if not os.path.exists(xlsx_path):
    xlsx_path = os.path.join(DATASET_DIR, 'Andhra_Pradesh_Master_College_Directory_2026.xlsx')

temp_load = os.path.join(DATASET_DIR, 'temp_cache_load.xlsx')
try:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
except Exception:
    import subprocess
    subprocess.run(["powershell", "-Command", f"Copy-Item '{xlsx_path}' '{temp_load}' -Force"], check=True)
    wb = openpyxl.load_workbook(temp_load, data_only=True)
    try:
        os.remove(temp_load)
    except Exception:
        pass

cache = {
    'sheets': {},
    'districts': [],
    'summary': {
        'total_colleges': 1832,
        'with_website': 1626,
        'without_website': 206,
        'with_email': 1662,
        'with_phone': 1768,
        'total_contacts': 2673
    }
}

for sheetname in wb.sheetnames:
    if sheetname == 'Executive Summary':
        continue
    ws = wb[sheetname]
    headers = [str(c or '').strip() for c in next(ws.iter_rows(max_row=1, values_only=True))]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(r):
            row_dict = dict(zip(headers, [str(v).strip() if v is not None else '' for v in r]))
            rows.append(row_dict)
    cache['sheets'][sheetname] = rows
    print(f"Sheet '{sheetname}': {len(rows)} rows")

# Extract unique districts from All Colleges AP Master
districts = set()
for r in cache['sheets'].get('All Colleges AP Master', []):
    d = r.get('District')
    if d and d not in ['0', 'None']:
        districts.add(str(d).strip())

cache['districts'] = sorted(districts)

output_json = os.path.join(DATASET_DIR, 'ap_master_dataset_cache.json')
with open(output_json, 'w', encoding='utf-8') as f:
    json.dump(cache, f, ensure_ascii=False)

# Also copy to public directory for direct fetch
os.makedirs(PUBLIC_DATASET_DIR, exist_ok=True)
public_json = os.path.join(PUBLIC_DATASET_DIR, 'ap_master_dataset_cache.json')
with open(public_json, 'w', encoding='utf-8') as f:
    json.dump(cache, f, ensure_ascii=False)

print(f"[OK] Cache saved to {output_json} ({os.path.getsize(output_json)} bytes)")
print(f"[OK] Synced to public at {public_json}")
